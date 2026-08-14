"""
UCL 財產/借用資料清洗與正規化腳本

用途：把三份格式不一致的資產表（A / B / 學校財產）與借用單（表單回應 1）
整理成統一的 Asset / BorrowRecord 兩張正規表，並標記每筆借用紀錄的配對狀態
（matched / pending_review / unlisted），保留原始值以利後續人工核對。

用法：
    python clean_ucl_data.py <財產清單.xlsx> <財產借用單.xlsx> <輸出.xlsx>
"""

import re
import sys
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from pathlib import Path

# ---------------------------------------------------------------------------
# 1. ID 正規化：去除 dash / 空白 / 全半形差異，只用來「比對」，不覆蓋原始值
# ---------------------------------------------------------------------------
def normalize_id_key(raw_id):
    """把 ID 轉成純比對用的 key。回傳 None 代表沒有可比對內容。"""
    if pd.isna(raw_id):
        return None
    s = str(raw_id).strip()
    if not s:
        return None
    s = s.replace("－", "-").replace("　", "").replace(" ", "")
    s = re.sub(r"[-]", "", s)
    return s.upper()


# ---------------------------------------------------------------------------
# 2. 讀取並合併三份資產表（A / B / 學校財產），統一成 Asset 正規表
# ---------------------------------------------------------------------------
def load_assets(asset_path):
    xls = pd.ExcelFile(asset_path)

    a = pd.read_excel(xls, sheet_name="A")
    b = pd.read_excel(xls, sheet_name="B")
    c = pd.read_excel(xls, sheet_name="學校財產")

    a = a.rename(columns={"實驗室編號": "asset_id"})
    a["id_type"] = "lab_code"

    b = b.rename(columns={"實驗室編號": "asset_id"})
    b["id_type"] = "lab_code"

    c = c.rename(columns={"學校財產編號": "asset_id"})
    c["id_type"] = "school_code"
    if "學校財產名稱" in c.columns:
        c = c.rename(columns={"學校財產名稱": "category_name"})
    else:
        c["category_name"] = None

    for df, src in [(a, "A"), (b, "B"), (c, "學校財產")]:
        df["source_sheet"] = src

    common_cols = {
        "asset_id": "asset_id",
        "id_type": "id_type",
        "品名": "name",
        "存放地點": "location",
        "數量": "qty",
        "狀態(正常、損壞、待報廢)": "status",
        "規格": "spec",
        "備註": "note",
        "source_sheet": "source_sheet",
    }

    frames = []
    for df in (a, b, c):
        cat = df["category_name"] if "category_name" in df.columns else None
        sub = pd.DataFrame({
            "asset_id": df.get("asset_id"),
            "id_type": df.get("id_type"),
            "name": df.get("品名"),
            "category_name": cat,
            "location": df.get("存放地點"),
            "qty": df.get("數量"),
            "status": df.get("狀態(正常、損壞、待報廢)"),
            "spec": df.get("規格"),
            "note": df.get("備註"),
            "source_sheet": df.get("source_sheet"),
        })
        frames.append(sub)

    assets = pd.concat(frames, ignore_index=True)
    assets["asset_id"] = assets["asset_id"].astype(str).str.strip()
    assets["id_key"] = assets["asset_id"].apply(normalize_id_key)

    # 標記重複 ID（同一 asset_id 出現多次），不自動合併，留給人工核對
    dup_mask = assets["asset_id"].duplicated(keep=False) & assets["asset_id"].notna()
    assets["is_duplicate_id"] = dup_mask

    return assets


# ---------------------------------------------------------------------------
# 3. 讀取並清洗借用單（只用「表單回應 1」，排除空白的 20221018 與髒亂的舊版借用單）
# ---------------------------------------------------------------------------
def load_borrows(borrow_path):
    xls = pd.ExcelFile(borrow_path)
    df = pd.read_excel(xls, sheet_name="表單回應 1")

    other_col = "如果你選擇了“其他”，請填此項"
    id_col = "設備編號or產編（如果有的話，請填寫）"

    def resolve_item_name(row):
        if row["借用器材"] == "其他" and pd.notna(row.get(other_col)):
            return str(row[other_col]).strip()
        return row["借用器材"]

    out = pd.DataFrame({
        "record_id": range(1, len(df) + 1),
        "timestamp": df["時間戳記"],
        "borrower_email": df["電子郵件地址"],
        "handler_name": df["經手人"],
        "item_name": df.apply(resolve_item_name, axis=1),
        "raw_asset_id": df[id_col].astype(str).where(df[id_col].notna(), None),
        "qty": df["數量（請寫純數字）"],
        "purpose": df["用途（自用、實驗用、上課用，其他用途不限）"],
        "use_location": df["使用地點（自家、教室名：C217等）"],
        "borrow_date": df["出借日期"],
        "return_date": df["歸還日期"],
        "note": df["備註欄"],
    })

    # 歸還日期空白 = 目前仍借出中，這是業務語意，不是缺值
    out["is_returned"] = out["return_date"].notna()

    # 數量缺值補 1，並標記為推測值，不悄悄補
    out["qty_is_estimated"] = out["qty"].isna()
    out["qty"] = out["qty"].fillna(1).astype(int)

    out["raw_id_key"] = out["raw_asset_id"].apply(normalize_id_key)

    return out


# ---------------------------------------------------------------------------
# 4. 借用紀錄 <-> 資產表配對，分成 matched / pending_review / unlisted 三桶
# ---------------------------------------------------------------------------
def match_borrows(borrows, assets):
    assets_by_key = {}
    for _, row in assets.iterrows():
        if row["id_key"]:
            assets_by_key.setdefault(row["id_key"], row)

    names_by_key = {}
    for _, row in assets.iterrows():
        if pd.notna(row["name"]):
            names_by_key.setdefault(str(row["name"]).strip().lower(), row)

    statuses, matched_ids, matched_names, reasons = [], [], [], []

    for _, row in borrows.iterrows():
        raw_id = row["raw_asset_id"]
        raw_key = row["raw_id_key"]
        status, m_id, m_name, reason = "unlisted", None, None, ""

        # 注意：raw_key 可能是 NaN(float)，NaN 在 Python 是 truthy，
        # 用 pd.notna() 而非直接 `if raw_key:` 判斷，避免誤判
        if pd.notna(raw_key):
            if raw_key in assets_by_key:
                status = "matched"
                m_id = assets_by_key[raw_key]["asset_id"]
                m_name = assets_by_key[raw_key]["name"]
            else:
                # 處理逗號分隔多值，如 "4050202-01-5000016,13號"
                parts = re.split(r"[,，]", str(raw_id))
                found = None
                for p in parts:
                    k = normalize_id_key(p)
                    if k and k in assets_by_key:
                        found = assets_by_key[k]
                        break
                if found is not None:
                    status = "matched"
                    m_id = found["asset_id"]
                    m_name = found["name"]
                elif raw_key.isdigit():
                    status = "pending_review"
                    reason = "純數字ID，來源系統不明，需人工核對"
                else:
                    status = "pending_review"
                    reason = "ID格式正確但資產表查無此ID，需確認資產表是否完整"

        if status == "unlisted":
            key = str(row["item_name"]).strip().lower()
            if key in names_by_key:
                status = "matched"
                m_id = names_by_key[key]["asset_id"]
                m_name = names_by_key[key]["name"]
            elif row["item_name"] != "其他":
                reason = "無資產編號，品名亦查無對應，可能為未列管/私人設備"
            else:
                reason = "無資產編號、無有效品名"

        statuses.append(status)
        matched_ids.append(m_id)
        matched_names.append(m_name)
        reasons.append(reason)

    borrows = borrows.copy()
    borrows["match_status"] = statuses
    borrows["matched_asset_id"] = matched_ids
    borrows["matched_asset_name"] = matched_names
    borrows["match_note"] = reasons
    return borrows


# ---------------------------------------------------------------------------
# 5. 輸出成正規化的 Excel（多分頁：資產清單 / 借用紀錄 / 待核對清單 / 摘要）
# ---------------------------------------------------------------------------
def export_regularized(assets, borrows, output_path):
    assets_out = assets[[
        "asset_id", "id_type", "name", "category_name", "location", "qty",
        "status", "spec", "note", "source_sheet", "is_duplicate_id",
    ]].rename(columns={
        "asset_id": "資產編號", "id_type": "編號類型", "name": "品名",
        "category_name": "財產類別", "location": "存放地點", "qty": "數量",
        "status": "狀態", "spec": "規格", "note": "備註",
        "source_sheet": "來源分頁", "is_duplicate_id": "重複ID警示",
    })

    borrows_out = borrows[[
        "record_id", "timestamp", "borrower_email", "handler_name", "item_name",
        "raw_asset_id", "matched_asset_id", "matched_asset_name", "match_status",
        "match_note", "qty", "qty_is_estimated", "purpose", "use_location",
        "borrow_date", "return_date", "is_returned", "note",
    ]].rename(columns={
        "record_id": "紀錄編號", "timestamp": "時間戳記", "borrower_email": "借用人Email",
        "handler_name": "經手人", "item_name": "品名（已還原其他欄）",
        "raw_asset_id": "原填設備編號", "matched_asset_id": "配對到的資產編號",
        "matched_asset_name": "配對到的資產品名", "match_status": "配對狀態",
        "match_note": "備註說明", "qty": "數量", "qty_is_estimated": "數量為推測值",
        "purpose": "用途", "use_location": "使用地點", "borrow_date": "出借日期",
        "return_date": "歸還日期", "is_returned": "已歸還",
        "note": "原始備註",
    })

    pending = borrows_out[borrows_out["配對狀態"] == "pending_review"]
    unlisted = borrows_out[borrows_out["配對狀態"] == "unlisted"]

    summary = pd.DataFrame({
        "項目": [
            "資產總筆數", "唯一資產ID數", "重複ID筆數",
            "借用紀錄總筆數", "配對成功 (matched)", "待人工核對 (pending_review)",
            "未列管/查無對應 (unlisted)",
        ],
        "數值": [
            len(assets_out), assets["asset_id"].nunique(), int(assets["is_duplicate_id"].sum()),
            len(borrows_out),
            int((borrows_out["配對狀態"] == "matched").sum()),
            int((borrows_out["配對狀態"] == "pending_review").sum()),
            int((borrows_out["配對狀態"] == "unlisted").sum()),
        ],
    })

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="摘要", index=False)
        assets_out.to_excel(writer, sheet_name="資產清單_整合", index=False)
        borrows_out.to_excel(writer, sheet_name="借用紀錄_整合", index=False)
        pending.to_excel(writer, sheet_name="待核對清單", index=False)
        unlisted.to_excel(writer, sheet_name="未列管清單", index=False)

    _apply_formatting(output_path)


def _apply_formatting(path):
    wb = load_workbook(path)
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    body_font = Font(name="Arial")

    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
        ws.freeze_panes = "A2"
        for col_cells in ws.columns:
            length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
            col_letter = get_column_letter(col_cells[0].column)
            ws.column_dimensions[col_letter].width = min(max(length + 2, 10), 45)

    wb.save(path)


# ---------------------------------------------------------------------------
if __name__ == "__main__":
# 取得目前腳本所在的絕對路徑
    base_dir = Path(__file__).resolve().parent
    
    # 指定輸出目標資料夾：./data/raw
    output_dir = base_dir / "data" / "raw"
    output_dir.mkdir(parents=True, exist_ok=True)  # 如果資料夾不存在則自動建立

    # 處理命令列參數（假設命令列只傳入檔案名稱，或容許完整路徑）
    asset_path = sys.argv[1]
    borrow_path = sys.argv[2]
    
    # 如果使用者沒有在命令列指定完整輸出路徑，就預設存到 data/raw/
    if len(sys.argv) > 3:
        output_filename = sys.argv[3]
    else:
        output_filename = "cleaned_output.xlsx"  # 預設檔名
        
    output_path = output_dir / output_filename

    # 執行清洗流程
    assets = load_assets(asset_path)
    borrows = load_borrows(borrow_path)
    borrows = match_borrows(borrows, assets)
    export_regularized(assets, borrows, str(output_path))

    print(f"已輸出：{output_path.resolve()}")
    print(borrows["match_status"].value_counts())